"""Create or update one person's account and import their CVs from a folder.

Example (run from the project folder):
  python manage.py setup_candidate --username anish --full-name "Anish Shrestha" \
      --email you@example.com --admin --cv-folder "D:/path/Job Positions"
Each CV is imported into THIS person's profile only.
"""
import re
from pathlib import Path

from django.contrib.auth.models import User
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from agent.models import AuditLog, Candidate, CVDocument


def read_cv_file(path):
    if path.suffix.lower() == ".docx":
        import docx

        return "\n".join(p.text for p in docx.Document(str(path)).paragraphs).strip()
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader

        return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages).strip()
    return path.read_text(encoding="utf-8", errors="ignore").strip()


def clean_name(folder_name):
    return re.sub(r"^\d+\s*", "", folder_name).strip()


def load_index(folder):
    """Read an optional CV_Index.xlsx ('All CVs' sheet) -> {folder name: row dict}."""
    index_file = next(iter(folder.glob("CV_Index.xlsx")), None)
    if not index_file:
        return {}
    import openpyxl

    wb = openpyxl.load_workbook(index_file, data_only=True)
    ws = wb["All CVs"] if "All CVs" in wb.sheetnames else wb.worksheets[-1]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}
    header = [str(h or "").strip() for h in rows[0]]
    result = {}
    for r in rows[1:]:
        row = dict(zip(header, r))
        folder_col = str(row.get("Folder (in D:\\JobSearchAgent)") or row.get("Folder") or "")
        if folder_col:
            result[re.split(r"[\\/]", folder_col.rstrip("\\/"))[-1].lower()] = row
    return result


class Command(BaseCommand):
    help = "Create/update a person's login and profile, and import their CV files (.docx, .pdf, .txt)."

    def add_arguments(self, p):
        p.add_argument("--username", required=True)
        p.add_argument("--full-name", required=True)
        p.add_argument("--email", default="")
        p.add_argument("--phone", default="")
        p.add_argument("--location", default="Auckland")
        p.add_argument("--job-types", default="", help="Comma separated job types for the daily search.")
        p.add_argument("--availability", default="")
        p.add_argument("--notes", default="", help="Extra true facts for the AI.")
        p.add_argument("--admin", action="store_true", help="Make this login an admin.")
        p.add_argument("--cv-folder", help="Folder with CV files (searched in sub-folders too).")

    @transaction.atomic
    def handle(self, *args, **o):
        user, created = User.objects.get_or_create(username=o["username"], defaults={"email": o["email"]})
        if created:
            user.set_unusable_password()
            self.stdout.write(f"Created login '{user.username}' (set a password with set_password.bat).")
        user.email = o["email"] or user.email
        user.first_name = o["full_name"][:150]
        if o["admin"]:
            user.is_staff = user.is_superuser = True
        user.save()

        candidate, _ = Candidate.objects.get_or_create(user=user, defaults={"full_name": o["full_name"]})
        for field, key in [("full_name", "full_name"), ("email", "email"), ("phone", "phone"),
                           ("location", "location"), ("job_types", "job_types"),
                           ("availability", "availability"), ("extra_notes", "notes")]:
            if o[key]:
                setattr(candidate, field, o[key])
        candidate.save()

        imported = 0
        if o.get("cv_folder"):
            folder = Path(o["cv_folder"])
            if not folder.is_dir():
                raise CommandError(f"Folder not found: {folder}")
            index = load_index(folder)
            files = sorted(folder.rglob("*.docx")) + sorted(folder.rglob("*.pdf")) + sorted(folder.rglob("*.txt"))
            done_dirs = set()
            for f in files:
                if f.name.startswith("~$") or f.parent in done_dirs:
                    continue  # one CV per folder; Word (.docx) preferred over PDF
                text = read_cv_file(f)
                if len(text) < 100:
                    continue
                done_dirs.add(f.parent)
                row = index.get(f.parent.name.lower(), {})
                title = str(row.get("Job Position") or clean_name(f.parent.name)) + " CV"
                parts = f.relative_to(folder).parts
                category = str(row.get("Category") or (clean_name(parts[-3]) if len(parts) >= 3 else ""))
                CVDocument.objects.update_or_create(
                    candidate=candidate, title=title[:120],
                    defaults={"text": text, "category": category[:120],
                              "best_for": str(row.get("Best For") or "")[:255],
                              "source_path": str(f.relative_to(folder))[:500]})
                imported += 1
        AuditLog.log("candidate_setup", f"{candidate.full_name}: {imported} CVs imported", candidate=candidate)
        self.stdout.write(self.style.SUCCESS(
            f"{candidate.full_name}: profile saved, {imported} CVs imported, {candidate.cvs.count()} CVs in total."))
